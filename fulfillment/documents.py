"""Generate customer-specific initial versions from production clinical masters."""

from __future__ import annotations

from datetime import datetime, timezone
from html import escape
from io import BytesIO
import re
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from fulfillment.neuromodulators import DOCUMENTS as NEUROMODULATOR_DOCUMENTS
from fulfillment.neuromodulators import state_language as neuromodulator_state_language


INITIAL_VERSION_LABEL = "INITIAL VERSION — Prepared for Qualified Provider Review"
PRODUCTION_MASTER_TREATMENTS = {"Neuromodulators"}
NAVY = RGBColor(8, 40, 93)
TEAL = RGBColor(6, 153, 159)


def _clean(value: object, fallback: str = "Not provided") -> str:
    text = " ".join(str(value or "").split())
    return text or fallback


def _first(mapping: dict, *keys: str, fallback: str = "Not provided") -> str:
    for key in keys:
        value = mapping.get(key)
        if value is not None and str(value).strip():
            return _clean(value)
    return fallback


def validate_production_treatments(treatments: list[str]) -> None:
    unsupported = [name for name in treatments if name not in PRODUCTION_MASTER_TREATMENTS]
    if unsupported:
        joined = ", ".join(unsupported)
        raise ValueError(
            f"The production master package is not yet connected to automated fulfillment for: {joined}. "
            "Keep this order in scope review until its approved clinical master is connected."
        )


def _context(order: dict, treatment: str) -> dict[str, str]:
    clinic = order.get("clinic") or {}
    oversight = order.get("oversight") or {}
    detail = (order.get("details") or {}).get(treatment) or {}
    state = _first(clinic, "state", "primaryState", "primary_state")
    clinic_name = _first(clinic, "name", "clinicName", "clinic_name", "businessName", "practiceName")
    address = _first(clinic, "address", "clinicAddress", "streetAddress", fallback="Address to be completed by clinic")
    phone = _first(clinic, "phone", "clinicPhone", "telephone", fallback="Clinic to complete before patient use")
    director = _first(oversight, "director", "medicalDirector", "reviewer", fallback="Qualified medical director / supervising provider to complete")
    providers = ", ".join(order.get("providers") or []) or "Authorized clinician roles to be confirmed"
    product = _first(detail, "product", "medication", "device", fallback="Clinic-selected authorized botulinum toxin type A product")
    dose = _first(detail, "dose", "dosePlan", "settings", fallback="Patient-specific dose to be authorized by the qualified prescriber using the selected product's current prescribing information")
    areas = _first(detail, "areas", "plannedAreas", "method", fallback="Patient-specific treatment areas to be documented before treatment")
    indication = _first(detail, "indication", "goal", fallback="Cosmetic neuromodulator treatment based on individualized assessment")
    state_text = neuromodulator_state_language(state)
    return {
        "[[CLINIC_NAME]]": clinic_name,
        "[[STATE]]": state,
        "[[STATE / ADDRESS]]": f"{state} / {address}",
        "[[MEDICAL_DIRECTOR_NAME, CREDENTIALS]]": director,
        "[[NAME, CREDENTIALS, LICENSE]]": director,
        "[[AUTHORIZED_CREDENTIALS]]": providers,
        "[[STATE-SPECIFIC AUTHORIZED ROLES]]": providers,
        "[[AUTHORIZED_PRODUCTS]]": product,
        "[[PRODUCT]]": product,
        "[[PRODUCT_1]]": product,
        "[[PRODUCT_2]]": "Not selected",
        "[[PRODUCT_3]]": "Not selected",
        "[[STATE_SPECIFIC_NEUROMODULATOR_REQUIREMENTS]]": state_text,
        "[[STATE_SPECIFIC_REQUIREMENT]]": state_text,
        "[[CLINIC_PHONE]]": phone,
        "[[EMERGENCY_PHONE]]": phone,
        "[[EMERGENCY_INSTRUCTIONS]]": f"For emergency symptoms activate EMS/911. Clinic contact: {phone}.",
        "[[MEDICAL_DIRECTOR_CONTACT]]": director,
        "[[INDICATION]]": indication,
        "[[PLANNED_AREAS]]": areas,
        "[[DOSE_PLAN]]": dose,
        "[[FOLLOW_UP_WINDOW]]": "Product- and area-appropriate follow-up; many aesthetic practices assess peak effect around 2 weeks.",
        "[[EFFECTIVE_DATE]]": "To be assigned at qualified-provider approval",
        "[[REVIEW_DATE]]": "To be assigned at qualified-provider approval",
        "[[GENERIC]]": "Product-specific; verify current prescribing information",
        "[[STRENGTH]]": "Product-specific; verify current prescribing information",
        "[[PI-BASED PREPARATION]]": "Follow current prescribing information for the selected product",
        "[[AREAS]]": areas,
    }


def _customize(text: str, replacements: dict[str, str]) -> str:
    for key, value in replacements.items():
        text = text.replace(key, value)
    # Never expose unresolved template tokens as if they were completed clinical content.
    return re.sub(r"\[\[[^\]]+\]\]", "To be completed during qualified-provider review", text)


def _is_heading(line: str) -> bool:
    if re.match(r"^\d+\.\s+", line):
        return True
    headings = {
        "Treatment being considered", "How the treatment works", "Potential benefits",
        "Common and expected effects", "Important risks", "Pregnancy, breastfeeding and medical conditions",
        "Medications and prior toxin treatment", "Approved and off-label use", "Alternatives",
        "Aftercare and follow-up", "Photography and privacy", "Acknowledgment", "Signature",
        "Before your appointment", "After treatment - what is usually expected", "Protect the treatment result",
        "Call the clinic promptly for", "Seek urgent/emergency medical care for", "Opening / room readiness",
        "Medication safety", "Before patient enters / before injection", "Emergency readiness check", "Closeout",
        "Do not treat until resolved or specifically cleared", "Pregnancy and lactation", "Assessment elements",
        "Clinic authorization field",
    }
    return line in headings


def _shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def _make_docx(title: str, text: str, order: dict) -> bytes:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.62)
    section.bottom_margin = Inches(0.62)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9.5)

    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = brand.add_run("MYCLINIC")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = NAVY
    run = brand.add_run("PROTOCOLS")
    run.bold = True
    run.font.size = Pt(16)
    run.font.color.rgb = TEAL

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = NAVY

    status = doc.add_paragraph()
    status.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = status.add_run(INITIAL_VERSION_LABEL)
    run.bold = True
    run.font.size = Pt(10)
    run.font.color.rgb = TEAL

    clinic = order.get("clinic") or {}
    meta = doc.add_table(rows=2, cols=2)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.style = "Table Grid"
    values = [
        ("Clinic", _first(clinic, "name", "clinicName", "clinic_name", "businessName", "practiceName")),
        ("State", _first(clinic, "state", "primaryState", "primary_state")),
        ("Order", _clean(order.get("orderReference"))),
        ("Generated", datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC")),
    ]
    for idx, (label, value) in enumerate(values):
        cell = meta.cell(idx // 2, idx % 2)
        _shade_cell(cell, "EAF8F8")
        cell.text = f"{label}: {value}"
        for r in cell.paragraphs[0].runs:
            r.font.size = Pt(8.5)
            r.font.color.rgb = NAVY
    doc.add_paragraph()

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    # The first line is the internal source marker; the title is already rendered above.
    if lines and lines[0] == "CLINICAL MASTER TEMPLATE":
        lines = lines[1:]
    # Skip duplicate source title/subtitle lines before the first substantive section.
    while lines and not _is_heading(lines[0]) and len(lines) > 1:
        if lines[0].lower() in title.lower() or title.lower() in lines[0].lower() or "•" in lines[0]:
            lines.pop(0)
        else:
            break

    for line in lines:
        if _is_heading(line):
            heading = doc.add_heading(line, level=1)
            for r in heading.runs:
                r.font.color.rgb = NAVY
                r.font.size = Pt(13)
        elif line.startswith("[ ]"):
            doc.add_paragraph(line)
        else:
            doc.add_paragraph(line)

    footer_note = doc.add_paragraph()
    footer_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_note.add_run(
        "MyClinicProtocols Initial Version • Final clinical approval remains with the purchasing clinic’s qualified medical director or supervising provider."
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(95, 113, 141)

    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()


def _make_pdf(title: str, text: str, order: dict) -> bytes:
    stream = BytesIO()
    doc = SimpleDocTemplate(
        stream,
        pagesize=LETTER,
        rightMargin=0.62 * inch,
        leftMargin=0.62 * inch,
        topMargin=0.55 * inch,
        bottomMargin=0.55 * inch,
        title=title,
        author="MyClinicProtocols",
    )
    styles = getSampleStyleSheet()
    brand_style = ParagraphStyle("Brand", parent=styles["Heading2"], textColor=colors.HexColor("#08285D"), alignment=TA_CENTER, fontSize=13, leading=16, spaceAfter=5)
    title_style = ParagraphStyle("TitleMYCP", parent=styles["Title"], textColor=colors.HexColor("#08285D"), alignment=TA_CENTER, fontSize=18, leading=22, spaceAfter=7)
    status_style = ParagraphStyle("Status", parent=styles["Normal"], textColor=colors.HexColor("#06999F"), alignment=TA_CENTER, fontName="Helvetica-Bold", fontSize=9, leading=12, spaceAfter=12)
    heading_style = ParagraphStyle("H1MYCP", parent=styles["Heading2"], textColor=colors.HexColor("#08285D"), fontSize=12, leading=15, spaceBefore=8, spaceAfter=4)
    body_style = ParagraphStyle("BodyMYCP", parent=styles["BodyText"], textColor=colors.HexColor("#263D5D"), fontSize=8.4, leading=11.2, spaceAfter=4)
    meta_style = ParagraphStyle("Meta", parent=body_style, backColor=colors.HexColor("#EAF8F8"), borderPadding=7, spaceAfter=9)
    small_style = ParagraphStyle("Small", parent=body_style, textColor=colors.HexColor("#65758C"), alignment=TA_CENTER, fontSize=7.5, leading=10, spaceBefore=9)

    clinic = order.get("clinic") or {}
    story = [
        Paragraph("<b>MYCLINIC</b><font color='#06999F'><b>PROTOCOLS</b></font>", brand_style),
        Paragraph(escape(title), title_style),
        Paragraph(escape(INITIAL_VERSION_LABEL), status_style),
        Paragraph(
            f"<b>Clinic:</b> {escape(_first(clinic, 'name', 'clinicName', 'clinic_name', 'businessName', 'practiceName'))} &nbsp;&nbsp; "
            f"<b>State:</b> {escape(_first(clinic, 'state', 'primaryState', 'primary_state'))}<br/>"
            f"<b>Order:</b> {escape(_clean(order.get('orderReference')))} &nbsp;&nbsp; "
            f"<b>Generated:</b> {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}",
            meta_style,
        ),
    ]
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if lines and lines[0] == "CLINICAL MASTER TEMPLATE":
        lines = lines[1:]
    while lines and not _is_heading(lines[0]) and len(lines) > 1:
        if lines[0].lower() in title.lower() or title.lower() in lines[0].lower() or "•" in lines[0]:
            lines.pop(0)
        else:
            break
    for line in lines:
        safe = escape(line).replace("•", "&bull;")
        if _is_heading(line):
            story.append(Paragraph(safe, heading_style))
        elif line.startswith("[ ]"):
            story.append(Paragraph("&#9633; " + escape(line[3:].strip()), body_style))
        else:
            story.append(Paragraph(safe, body_style))
    story.extend([
        Spacer(1, 6),
        Paragraph(
            "MyClinicProtocols Initial Version • Final clinical approval remains with the purchasing clinic’s qualified medical director or supervising provider.",
            small_style,
        ),
    ])
    doc.build(story)
    return stream.getvalue()


def _neuromodulator_files(order: dict) -> list[tuple[str, bytes, bytes]]:
    replacements = _context(order, "Neuromodulators")
    package_type = order.get("package", "protocol")
    source_docs = NEUROMODULATOR_DOCUMENTS if package_type == "complete" else NEUROMODULATOR_DOCUMENTS[:1]
    files: list[tuple[str, bytes, bytes]] = []
    for slug, title, source in source_docs:
        text = _customize(source, replacements)
        files.append((slug, _make_docx(title, text, order), _make_pdf(title, text, order)))
    return files


def make_package(order: dict) -> bytes:
    treatments = order.get("treatments") or []
    if not treatments:
        raise ValueError("At least one treatment is required")
    validate_production_treatments(treatments)

    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        for treatment in treatments:
            if treatment == "Neuromodulators":
                for slug, docx_bytes, pdf_bytes in _neuromodulator_files(order):
                    archive.writestr(f"{slug}-INITIAL-VERSION.docx", docx_bytes)
                    archive.writestr(f"{slug}-INITIAL-VERSION.pdf", pdf_bytes)
        archive.writestr(
            "READ-ME.txt",
            "MyClinicProtocols — Initial Version\n\n"
            "This package was generated from the production clinical master package and customized with the clinic, provider, product, and applicable state information supplied with the order.\n\n"
            "Final clinical approval remains with the purchasing clinic’s appropriately qualified medical director or supervising provider. MyClinicProtocols RN quality review follows the Initial Version and is normally completed within 1–2 hours, with up to 24 hours for larger or more complex document sets.\n"
            "Up to two consolidated revision rounds may be requested within 14 calendar days of the RN-reviewed delivery; revision requests normally take 3–5 business days.\n",
        )
    return stream.getvalue()
